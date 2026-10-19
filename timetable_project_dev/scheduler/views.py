import csv
import io
import re
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db.models import Q
from .models import Teacher, Student, Subject, Room, TeacherUnavailability, LessonRequirement, ScheduleItem
from .solver import TimetableCSPSolver

DAYS_ORDER = [
    ("Monday", "Понеділок"),
    ("Tuesday", "Вівторок"),
    ("Wednesday", "Середа"),
    ("Thursday", "Четвер"),
    ("Friday", "П'ятниця")
]

def normalize_subject_title(title):
    if not title:
        return ""
    t = str(title).strip()
    m = re.search(r'=HYPERLINK\([^,]+,\s*"([^"]+)"\)', t, re.IGNORECASE)
    if m:
        t = m.group(1).strip()
    t = re.sub(r'\s+', ' ', t).strip()
    t = re.sub(r'\s+([.,])', r'\1', t)
    return t

def are_similar_subjects(s1, s2):
    n1 = normalize_subject_title(s1).lower()
    n2 = normalize_subject_title(s2).lower()
    if n1 == n2:
        return True
    c1 = re.sub(r'[\(\)\.\,\s\-]+', '', n1)
    c2 = re.sub(r'[\(\)\.\,\s\-]+', '', n2)
    return c1 == c2

def detect_faculty_name(filename, sheet):
    for r in range(1, 8):
        for c in range(1, 6):
            val = str(sheet.cell(row=r, column=c).value or '').strip()
            if any(k in val.lower() for k in ['факультет', 'інститут', 'навчально-науковий']):
                return val
    base = os.path.splitext(filename)[0]
    parts = re.split(r'[\s_]+', base)
    if parts and len(parts[0].strip()) > 1:
        return parts[0].strip().upper()
    return "Факультет"

def guess_course_from_group(group_name):
    g = str(group_name).strip()
    if g.upper().startswith('М'):
        return 5
    digits = re.findall(r'\d', g)
    if digits:
        last_d = digits[-1]
        mapping = {'6': 1, '5': 2, '4': 3, '3': 4, '2': 5, '1': 1}
        return mapping.get(last_d, 1)
    return 1

def schedule_dashboard(request):
    current_week = int(request.GET.get('week', 1))
    selected_group_filter = request.GET.get('filter_group', '').strip()
    if selected_group_filter:
        request.session['saved_group_filter'] = selected_group_filter
    elif 'saved_group_filter' in request.session:
        selected_group_filter = request.session['saved_group_filter']

    if 'filter_group' in request.GET and not request.GET.get('filter_group'):
        if 'saved_group_filter' in request.session:
            del request.session['saved_group_filter']
        selected_group_filter = ''

    edit_item_id = request.GET.get('edit_item')
    editing_item = ScheduleItem.objects.filter(id=edit_item_id, week=current_week).first() if edit_item_id else None

    edit_unavail_id = request.GET.get('edit_unavail')
    editing_unavail = TeacherUnavailability.objects.filter(id=edit_unavail_id).first() if edit_unavail_id else None

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "add_teacher":
            name = request.POST.get("name", "").strip()
            if name:
                Teacher.objects.create(name=name)
                messages.success(request, f"Викладача «{name}» додано!")
            return redirect(f"/?week={current_week}")
        elif action == "add_subject":
            title = request.POST.get("title", "").strip()
            is_elective = bool(request.POST.get("is_elective"))
            level = request.POST.get("level", "bachelor")
            if title:
                Subject.objects.create(title=title, is_elective=is_elective, level=level)
                messages.success(request, f"Дисципліну «{title}» додано!")
            return redirect(f"/?week={current_week}")
        elif action == "add_room":
            number = request.POST.get("number", "").strip()
            if number:
                Room.objects.get_or_create(number=number)
                messages.success(request, f"Аудиторію {number} додано!")
            return redirect(f"/?week={current_week}")

    items_qs = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')

    all_groups_set = set(Student.objects.values_list('group_name', flat=True))
    for it in items_qs:
        for g in it.groups.split(','):
            g_clean = g.strip()
            if g_clean:
                all_groups_set.add(g_clean)
    all_groups_list = sorted(list(all_groups_set))

    grid = {p: {d[0]: [] for d in DAYS_ORDER} for p in range(1, 9)}
    for it in items_qs:
        if selected_group_filter:
            item_groups = [g.strip().lower() for g in it.groups.split(',')]
            if not any(selected_group_filter.lower() in ig for ig in item_groups) and it.groups:
                continue
        if it.period in grid and it.day in grid[it.period]:
            grid[it.period][it.day].append(it)

    context = {
        'current_week': current_week,
        'selected_group_filter': selected_group_filter,
        'all_groups_list': all_groups_list,
        'editing_item': editing_item,
        'editing_unavail': editing_unavail,
        'teachers': Teacher.objects.all(),
        'subjects': Subject.objects.all(),
        'rooms': Room.objects.all(),
        'unavailabilities': TeacherUnavailability.objects.select_related('teacher').all(),
        'requirements': LessonRequirement.objects.filter(week=current_week).select_related('teacher', 'subject'),
        'days_order': DAYS_ORDER,
        'periods': range(1, 9),
        'grid': grid
    }
    return render(request, "schedule.html", context)

def import_schedule_view(request):
    if request.method == "POST":
        file = request.FILES.get("schedule_file")
        if not file:
            messages.error(request, "Оберіть файл для завантаження!")
            return redirect("import_schedule")

        filename = file.name
        imported_count = 0
        try:
            if filename.lower().endswith(".xlsx"):
                wb = openpyxl.load_workbook(file, data_only=False)
                sheet = wb.active
                detected_faculty = detect_faculty_name(filename, sheet)

                week_num = 1
                sheet_title = sheet.title.lower()
                fn_lower = filename.lower()
                if any(x in sheet_title or x in fn_lower for x in ['2 тиждень', '2тиждень', '2_тиждень', 'друг', '2сем']):
                    week_num = 2

                group_row = None
                for r in range(1, 20):
                    c1 = str(sheet.cell(row=r, column=1).value or '').strip().lower()
                    c2 = str(sheet.cell(row=r, column=2).value or '').strip().lower()
                    if 'дні' in c1 or 'пари' in c2 or 'дни' in c1 or 'пары' in c2:
                        group_row = r + 1
                        break
                if not group_row:
                    group_row = 8

                groups_by_col = {}
                for c in range(3, sheet.max_column + 1):
                    g_val = sheet.cell(row=group_row, column=c).value
                    if g_val:
                        clean_g = str(g_val).strip()
                        if len(clean_g) > 0 and clean_g.lower() not in ['дні', 'пари', 'none', 'пара', 'день']:
                            groups_by_col[c] = clean_g

                for g_name in set(groups_by_col.values()):
                    crs = guess_course_from_group(g_name)
                    Student.objects.update_or_create(
                        group_name=g_name,
                        defaults={'full_name': f"Студент {g_name}", 'faculty': detected_faculty, 'course': crs}
                    )

                days_map = {
                    'понеділок': 'Monday', 'вівторок': 'Tuesday', 'середа': 'Wednesday',
                    'четвер': 'Thursday', "п'ятниця": 'Friday', 'пятниця': 'Friday',
                    'понедельник': 'Monday', 'вторник': 'Tuesday', 'среда': 'Wednesday',
                    'четверг': 'Thursday', 'пятница': 'Friday', 'субота': 'Saturday'
                }

                row_day_period = {}
                current_day = None
                for r in range(group_row + 1, sheet.max_row + 1):
                    d_val = sheet.cell(row=r, column=1).value
                    if d_val:
                        d_str = str(d_val).strip().lower()
                        for k, v in days_map.items():
                            if k in d_str:
                                current_day = v
                                break
                    p_val = sheet.cell(row=r, column=2).value
                    if p_val is not None:
                        try:
                            period_num = int(float(str(p_val).strip()))
                            if current_day:
                                row_day_period[r] = (current_day, period_num)
                        except:
                            pass

                default_teacher, _ = Teacher.objects.get_or_create(name=f"Викладач ({detected_faculty})")

                for r, (day, period) in row_day_period.items():
                    if day == 'Monday' and period > 7:
                        continue
                    slot_items = []
                    for col, group_name in groups_by_col.items():
                        cell_val = sheet.cell(row=r, column=col).value
                        if cell_val:
                            subj_clean = normalize_subject_title(cell_val)
                            if subj_clean and len(subj_clean) > 1 and not subj_clean.isdigit():
                                slot_items.append((group_name, subj_clean))

                    clusters = []
                    for g_name, s_title in slot_items:
                        found = False
                        for cl in clusters:
                            if are_similar_subjects(cl[0], s_title):
                                cl[1].add(g_name)
                                found = True
                                break
                        if not found:
                            clusters.append([s_title, {g_name}])

                    for canonical_title, group_set in clusters:
                        is_elective_val = any(w in canonical_title.upper() for w in ['ВК', 'ВИБІРК', 'ЗА ВИБОРОМ'])
                        is_master = any(g.upper().startswith('М') for g in group_set)
                        level_val = 'master' if is_master else 'bachelor'

                        subj_obj, _ = Subject.objects.get_or_create(
                            title=canonical_title,
                            defaults={'is_elective': is_elective_val, 'level': level_val}
                        )
                        if is_elective_val and not subj_obj.is_elective:
                            subj_obj.is_elective = True
                            subj_obj.level = level_val
                            subj_obj.save()

                        existing = ScheduleItem.objects.filter(
                            week=week_num, day=day, period=period, subject=subj_obj
                        ).first()

                        if existing:
                            existing_set = set([g.strip() for g in existing.groups.split(',') if g.strip()])
                            union_set = existing_set.union(group_set)
                            existing.groups = ", ".join(sorted(list(union_set)))
                            existing.is_stream = len(union_set) > 1
                            existing.save()
                        else:
                            ScheduleItem.objects.create(
                                week=week_num, day=day, period=period,
                                subject=subj_obj, teacher=default_teacher,
                                groups=", ".join(sorted(list(group_set))),
                                is_stream=len(group_set) > 1
                            )
                        imported_count += 1

                messages.success(request, f"Успішно імпортовано розклад «{detected_faculty}» ({week_num}-й тиждень): {imported_count} дисциплін.")
            else:
                messages.error(request, "Формат файлу має бути .xlsx!")
        except Exception as e:
            messages.error(request, f"Помилка обробки файлу: {str(e)}")
        return redirect("import_schedule")

    return render(request, "import.html")

def selective_bachelor_view(request):
    edit_id = request.GET.get('edit')
    editing_subject = Subject.objects.filter(id=edit_id, level='bachelor').first() if edit_id else None
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "save":
            subj_id = request.POST.get("subject_id")
            title = request.POST.get("title", "").strip()
            if subj_id:
                s = get_object_or_404(Subject, id=subj_id)
                s.title = title
                s.is_elective = True
                s.level = 'bachelor'
                s.save()
            else:
                Subject.objects.create(title=title, is_elective=True, level='bachelor')
            return redirect("selective_bachelor")
        elif action == "delete":
            Subject.objects.filter(id=request.POST.get("subject_id"), level='bachelor').delete()
            return redirect("selective_bachelor")
    return render(request, "selective_bachelor.html", {'subjects': Subject.objects.filter(is_elective=True, level='bachelor').order_by('title'), 'editing_subject': editing_subject})

def selective_master_view(request):
    edit_id = request.GET.get('edit')
    editing_subject = Subject.objects.filter(id=edit_id, level='master').first() if edit_id else None
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "save":
            subj_id = request.POST.get("subject_id")
            title = request.POST.get("title", "").strip()
            if subj_id:
                s = get_object_or_404(Subject, id=subj_id)
                s.title = title
                s.is_elective = True
                s.level = 'master'
                s.save()
            else:
                Subject.objects.create(title=title, is_elective=True, level='master')
            return redirect("selective_master")
        elif action == "delete":
            Subject.objects.filter(id=request.POST.get("subject_id"), level='master').delete()
            return redirect("selective_master")
    return render(request, "selective_master.html", {'subjects': Subject.objects.filter(is_elective=True, level='master').order_by('title'), 'editing_subject': editing_subject})

def public_student_timetable_view(request):
    faculties = Student.objects.exclude(faculty="").values_list('faculty', flat=True).distinct().order_by('faculty')
    selected_faculty = request.GET.get('faculty', '').strip()
    selected_course = request.GET.get('course', '').strip()
    selected_group = request.GET.get('group', '').strip()
    selected_student_id = request.GET.get('student_id', '').strip()
    current_week = int(request.GET.get('week', 1))

    selected_student = Student.objects.filter(id=selected_student_id).first() if selected_student_id else None

    group_to_show = ""
    if selected_student:
        group_to_show = selected_student.group_name
    elif selected_group:
        group_to_show = selected_group

    lessons = []
    if group_to_show:
        items = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')
        day_order = {"Monday": 1, "Tuesday": 2, "Wednesday": 3, "Thursday": 4, "Friday": 5}
        for it in items:
            groups_list = [g.strip().lower() for g in it.groups.split(',')]
            if any(group_to_show.strip().lower() in gl for gl in groups_list) or not it.groups:
                lessons.append(it)
        lessons.sort(key=lambda x: (day_order.get(x.day, 9), x.period))

    return render(request, "student_public_timetable.html", {
        'faculties': faculties,
        'selected_faculty': selected_faculty,
        'selected_course': selected_course,
        'selected_group': selected_group,
        'selected_student': selected_student,
        'selected_student_id': selected_student_id,
        'group_to_show': group_to_show,
        'lessons': lessons,
        'current_week': current_week,
    })

def public_teacher_timetable_view(request):
    departments = Teacher.objects.exclude(department="").values_list('department', flat=True).distinct().order_by('department')
    selected_dept = request.GET.get('department', '').strip()
    selected_teacher_id = request.GET.get('teacher_id', '').strip()
    current_week = int(request.GET.get('week', 1))

    selected_teacher = Teacher.objects.filter(id=selected_teacher_id).first() if selected_teacher_id else None
    lessons = []
    if selected_teacher:
        items = ScheduleItem.objects.filter(week=current_week, teacher=selected_teacher).select_related('teacher', 'subject', 'room')
        day_order = {"Monday": 1, "Tuesday": 2, "Wednesday": 3, "Thursday": 4, "Friday": 5}
        lessons = sorted(items, key=lambda x: (day_order.get(x.day, 9), x.period))

    return render(request, "teacher_public_timetable.html", {
        'departments': departments,
        'selected_dept': selected_dept,
        'selected_teacher': selected_teacher,
        'selected_teacher_id': selected_teacher_id,
        'lessons': lessons,
        'current_week': current_week,
    })

def student_filters_api(request):
    faculty = request.GET.get('faculty', '').strip()
    course = request.GET.get('course', '').strip()
    group = request.GET.get('group', '').strip()
    qs = Student.objects.all()
    if faculty:
        qs = qs.filter(faculty=faculty)
    courses = list(qs.values_list('course', flat=True).distinct().order_by('course'))
    if course:
        qs = qs.filter(course=int(course))
    groups = list(qs.values_list('group_name', flat=True).distinct().order_by('group_name'))
    students = []
    if group:
        qs = qs.filter(group_name=group)
        students = list(qs.values('id', 'full_name').order_by('full_name'))
    return JsonResponse({'courses': courses, 'groups': groups, 'students': students})

def teacher_filters_api(request):
    dept = request.GET.get('department', '').strip()
    qs = Teacher.objects.all()
    if dept:
        qs = qs.filter(department=dept)
    teachers = list(qs.values('id', 'name').order_by('name'))
    return JsonResponse({'teachers': teachers})

def students_view(request):
    edit_id = request.GET.get('edit')
    editing_student = Student.objects.filter(id=edit_id).first() if edit_id else None
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "save":
            student_id = request.POST.get("student_id")
            full_name = request.POST.get("full_name", "").strip()
            faculty = request.POST.get("faculty", "").strip()
            course = int(request.POST.get("course", 1))
            group_name = request.POST.get("group_name", "").strip()
            if student_id:
                s = get_object_or_404(Student, id=student_id)
                s.full_name = full_name
                s.faculty = faculty
                s.course = course
                s.group_name = group_name
                s.save()
            else:
                Student.objects.create(full_name=full_name, faculty=faculty, course=course, group_name=group_name)
            return redirect("students_list")
        elif action == "delete":
            Student.objects.filter(id=request.POST.get("student_id")).delete()
            return redirect("students_list")
    return render(request, "students.html", {
        'students': Student.objects.all().order_by("course", "group_name", "full_name"),
        'editing_student': editing_student,
        'faculties': Student.objects.exclude(faculty="").values_list('faculty', flat=True).distinct().order_by('faculty'),
        'groups': Student.objects.values_list('group_name', flat=True).distinct().order_by('group_name'),
    })

def student_schedule_view(request, student_id):
    student = get_object_or_404(Student, id=student_id)
    current_week = int(request.GET.get('week', 1))
    all_items = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')
    student_lessons = []
    for item in all_items:
        assigned_groups = [g.strip().lower() for g in item.groups.split(',')]
        if any(student.group_name.strip().lower() in ag for ag in assigned_groups) or not item.groups:
            student_lessons.append(item)
    day_map_order = {"Monday": 1, "Tuesday": 2, "Wednesday": 3, "Thursday": 4, "Friday": 5}
    student_lessons.sort(key=lambda x: (day_map_order.get(x.day, 9), x.period))
    return render(request, "student_schedule.html", {'student': student, 'current_week': current_week, 'lessons': student_lessons})

def teachers_view(request):
    edit_id = request.GET.get('edit')
    editing_teacher = Teacher.objects.filter(id=edit_id).first() if edit_id else None
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "save":
            teacher_id = request.POST.get("teacher_id")
            name = request.POST.get("name", "").strip()
            department = request.POST.get("department", "").strip()
            if teacher_id:
                t = get_object_or_404(Teacher, id=teacher_id)
                t.name = name
                t.department = department
                t.save()
            else:
                Teacher.objects.create(name=name, department=department)
            return redirect("teachers_list")
        elif action == "delete":
            Teacher.objects.filter(id=request.POST.get("teacher_id")).delete()
            return redirect("teachers_list")
    return render(request, "teachers.html", {'teachers': Teacher.objects.all().order_by("name"), 'editing_teacher': editing_teacher})

def export_schedule_excel(request):
    current_week = int(request.GET.get('week', 1))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{current_week} тиждень"

    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Segoe UI", size=10)
    bold_cell_font = Font(name="Segoe UI", size=11, bold=True)
    
    red_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    thin_border = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'), top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws.merge_cells('A1:F1')
    top_cell = ws['A1']
    top_cell.value = f"Розклад занять — {current_week} тиждень"
    top_cell.font = title_font
    top_cell.fill = red_fill
    top_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    headers = ["Пара", "Понеділок", "Вівторок", "Середа", "Четвер", "П'ятниця"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = red_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[2].height = 28

    items = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')
    grid = {p: {d[0]: [] for d in DAYS_ORDER} for p in range(1, 9)}
    for it in items:
        if it.period in grid and it.day in grid[it.period]:
            grid[it.period][it.day].append(it)

    days_keys = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    for p in range(1, 9):
        row_idx = p + 2
        ws.row_dimensions[row_idx].height = 85
        num_cell = ws.cell(row=row_idx, column=1, value=f"{p} пара")
        num_cell.font = bold_cell_font
        num_cell.fill = gray_fill
        num_cell.alignment = Alignment(horizontal="center", vertical="center")
        num_cell.border = thin_border

        for col_idx, d in enumerate(days_keys, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if d == "Monday" and p == 8:
                cell.value = "—"
                cell.fill = gray_fill
                cell.font = cell_font
            else:
                lessons_in_slot = grid[p].get(d, [])
                if lessons_in_slot:
                    lines = []
                    for item in lessons_in_slot:
                        room_str = f" (ауд. {item.room.number})" if item.room else ""
                        stream_prefix = "★ [ПОТІК] " if item.is_stream else ""
                        elective_prefix = "ВК " if item.subject.is_elective else ""
                        lines.append(f"{stream_prefix}{elective_prefix}{item.subject.title}\nВикладач: {item.teacher.name}{room_str}\nгр. {item.groups}")
                    cell.value = "\n---\n".join(lines)
                    cell.font = cell_font
                else:
                    cell.value = ""

    ws.column_dimensions['A'].width = 15
    for c in ['B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[c].width = 40

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="schedule_week_{current_week}.xlsx"'
    wb.save(response)
    return response

# Експорт розкладу для конкретної групи (бакалаври / магістри) з кнопкою на порталі
def export_group_schedule_excel(request):
    group_name = request.GET.get('group', '').strip()
    current_week = int(request.GET.get('week', 1))
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Група {group_name} ({current_week} тиж)"

    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Segoe UI", size=10)
    bold_cell_font = Font(name="Segoe UI", size=11, bold=True)
    
    red_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    thin_border = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'), top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws.merge_cells('A1:F1')
    top_cell = ws['A1']
    top_cell.value = f"Розклад для групи {group_name} — {current_week} тиждень"
    top_cell.font = title_font
    top_cell.fill = red_fill
    top_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    headers = ["Пара", "Понеділок", "Вівторок", "Середа", "Четвер", "П'ятниця"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = red_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[2].height = 28

    items = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')
    grid = {p: {d[0]: [] for d in DAYS_ORDER} for p in range(1, 9)}
    
    for it in items:
        groups_list = [g.strip().lower() for g in it.groups.split(',')]
        if any(group_name.lower() in gl for gl in groups_list) or not it.groups:
            if it.period in grid and it.day in grid[it.period]:
                grid[it.period][it.day].append(it)

    days_keys = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    for p in range(1, 9):
        row_idx = p + 2
        ws.row_dimensions[row_idx].height = 85
        num_cell = ws.cell(row=row_idx, column=1, value=f"{p} пара")
        num_cell.font = bold_cell_font
        num_cell.fill = gray_fill
        num_cell.alignment = Alignment(horizontal="center", vertical="center")
        num_cell.border = thin_border

        for col_idx, d in enumerate(days_keys, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if d == "Monday" and p == 8:
                cell.value = "—"
                cell.fill = gray_fill
                cell.font = cell_font
            else:
                lessons_in_slot = grid[p].get(d, [])
                if lessons_in_slot:
                    lines = []
                    for item in lessons_in_slot:
                        room_str = f" (ауд. {item.room.number})" if item.room else ""
                        stream_prefix = "★ [ПОТІК] " if item.is_stream else ""
                        elective_prefix = "ВК " if item.subject.is_elective else ""
                        lines.append(f"{stream_prefix}{elective_prefix}{item.subject.title}\nВикладач: {item.teacher.name}{room_str}\nгр. {item.groups}")
                    cell.value = "\n---\n".join(lines)
                    cell.font = cell_font
                else:
                    cell.value = ""

    ws.column_dimensions['A'].width = 15
    for c in ['B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[c].width = 40

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="schedule_group_{group_name}_week_{current_week}.xlsx"'
    wb.save(response)
    return response
