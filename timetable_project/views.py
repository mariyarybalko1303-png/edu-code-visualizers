import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render, redirect
from django.http import HttpResponse
from django.contrib import messages
from .models import Teacher, Subject, Room, TeacherUnavailability, LessonRequirement, ScheduleItem
from .solver import TimetableCSPSolver

DAYS_ORDER = [
    ("Monday", "Понеділок"),
    ("Tuesday", "Вівторок"),
    ("Wednesday", "Середа"),
    ("Thursday", "Четвер"),
    ("Friday", "П'ятниця")
]

def schedule_dashboard(request):
    current_week = int(request.GET.get('week', 1))
    edit_unavail_id = request.GET.get('edit_unavail')
    editing_unavail = None
    if edit_unavail_id:
        editing_unavail = TeacherUnavailability.objects.filter(id=edit_unavail_id).first()

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "add_teacher":
            name = request.POST.get("name", "").strip()
            if name:
                Teacher.objects.create(name=name)
                messages.success(request, f"Викладача «{name}» додано!")
            return redirect(f"/?week={current_week}")

        elif action == "add_subject":
            title = request.POST.get("title", "").strip()
            if title:
                Subject.objects.create(title=title)
                messages.success(request, f"Дисципліну «{title}» додано!")
            return redirect(f"/?week={current_week}")

        elif action == "add_room":
            number = request.POST.get("number", "").strip()
            if number:
                Room.objects.get_or_create(number=number)
                messages.success(request, f"Аудиторію {number} додано!")
            return redirect(f"/?week={current_week}")

        elif action == "manual_assign":
            day = request.POST.get("day")
            period = int(request.POST.get("period"))
            teacher_id = request.POST.get("teacher_id")
            subject_id = request.POST.get("subject_id")
            room_id = request.POST.get("room_id")
            groups = request.POST.get("groups", "").strip()
            is_stream = bool(request.POST.get("is_stream"))

            if day == "Monday" and period > 7:
                messages.error(request, "У понеділок не може бути 8-ї пари!")
                return redirect(f"/?week={current_week}")

            teacher = Teacher.objects.get(id=teacher_id)
            subject = Subject.objects.get(id=subject_id)
            room = Room.objects.get(id=room_id) if room_id else None

            if ScheduleItem.objects.filter(week=current_week, day=day, period=period, teacher=teacher).exists():
                messages.error(request, f"Колізія: {teacher.name} вже має пару в цей час!")
                return redirect(f"/?week={current_week}")

            if room and ScheduleItem.objects.filter(week=current_week, day=day, period=period, room=room).exists():
                messages.error(request, f"Колізія: аудиторія {room.number} уже зайнята на цій парі!")
                return redirect(f"/?week={current_week}")

            ScheduleItem.objects.update_or_create(
                week=current_week,
                day=day,
                period=period,
                defaults={
                    'teacher': teacher,
                    'subject': subject,
                    'room': room,
                    'groups': groups,
                    'is_stream': is_stream
                }
            )
            messages.success(request, f"Пару «{subject.title}» призначено в розклад!")
            return redirect(f"/?week={current_week}")

        elif action == "delete_lesson":
            item_id = request.POST.get("item_id")
            if item_id:
                item = ScheduleItem.objects.filter(id=item_id, week=current_week).first()
                if item:
                    subj_title = item.subject.title
                    item.delete()
                    messages.success(request, f"Пару «{subj_title}» видалено з розкладу!")
            return redirect(f"/?week={current_week}")

        elif action == "save_unavailability":
            unavail_id = request.POST.get("unavail_id")
            teacher_id = request.POST.get("teacher_id")
            day = request.POST.get("day")
            period = int(request.POST.get("period", 0))

            if unavail_id:
                un = TeacherUnavailability.objects.filter(id=unavail_id).first()
                if un:
                    un.teacher_id = teacher_id
                    un.day = day
                    un.period = period
                    un.save()
                    messages.success(request, f"Зайнятість викладача {un.teacher.name} оновлено!")
            else:
                TeacherUnavailability.objects.create(
                    teacher_id=teacher_id,
                    day=day,
                    period=period
                )
                messages.success(request, "Зайнятість викладача зафіксовано!")
            return redirect(f"/?week={current_week}")

        elif action == "delete_unavailability":
            unavail_id = request.POST.get("unavail_id")
            if unavail_id:
                un = TeacherUnavailability.objects.filter(id=unavail_id).first()
                if un:
                    t_name = un.teacher.name
                    un.delete()
                    messages.success(request, f"Обмеження для {t_name} видалено!")
            return redirect(f"/?week={current_week}")

        elif action == "add_requirement":
            teacher_id = request.POST.get("teacher_id")
            subject_id = request.POST.get("subject_id")
            groups = request.POST.get("groups", "").strip()
            is_stream = bool(request.POST.get("is_stream"))
            hours = int(request.POST.get("hours", 1))

            LessonRequirement.objects.create(
                week=current_week,
                teacher_id=teacher_id,
                subject_id=subject_id,
                groups=groups,
                is_stream=is_stream,
                hours_per_week=hours
            )
            messages.success(request, "Вимогу до навантаження додано!")
            return redirect(f"/?week={current_week}")

        elif action == "auto_generate":
            rooms = list(Room.objects.all())
            if not rooms:
                messages.error(request, "Додайте хоча б одну аудиторію!")
                return redirect(f"/?week={current_week}")

            requirements = list(LessonRequirement.objects.filter(week=current_week))
            if not requirements:
                messages.error(request, "Немає вимог навантаження для цього тижня!")
                return redirect(f"/?week={current_week}")

            unavailabilities = list(TeacherUnavailability.objects.all())

            solver = TimetableCSPSolver(
                week=current_week,
                teachers=Teacher.objects.all(),
                subjects=Subject.objects.all(),
                rooms=rooms,
                requirements=requirements,
                unavailabilities=unavailabilities
            )

            solution = solver.solve()
            if solution is None:
                messages.error(request, "Неможливо розставити пари без колізій з поточними обмеженнями!")
            else:
                ScheduleItem.objects.filter(week=current_week).delete()
                for item in solution:
                    ScheduleItem.objects.create(**item)
                messages.success(request, f"Розклад успішно згенеровано ({len(solution)} пар)!")

            return redirect(f"/?week={current_week}")

    items = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')
    grid = {p: {} for p in range(1, 9)}
    for it in items:
        grid[it.period][it.day] = it

    context = {
        'current_week': current_week,
        'editing_unavail': editing_unavail,
        'teachers': Teacher.objects.all(),
        'subjects': Subject.objects.all(),
        'rooms': Room.objects.all(),
        'unavailabilities': TeacherUnavailability.objects.select_related('teacher').all(),
        'requirements': LessonRequirement.objects.filter(week=current_week).select_related('teacher', 'subject'),
        'days_order': DAYS_ORDER,
        'periods': range(1, 9),
        'grid': grid
    }
    return render(request, "schedule.html", context)

def export_schedule_excel(request):
    current_week = int(request.GET.get('week', 1))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{current_week} тиждень"

    # Стилі оформлення
    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Segoe UI", size=10)
    bold_cell_font = Font(name="Segoe UI", size=10, bold=True)
    
    # 1 тиждень — червоний, 2 тиждень — зелений
    header_fill_color = "DC2626" if current_week == 1 else "16A34A"
    header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    card_fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    stream_fill = PatternFill(start_color="FEFCE8", end_color="FEFCE8", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Верхній заголовок
    ws.merge_cells('A1:F1')
    top_cell = ws['A1']
    top_cell.value = f"Розклад занять — {current_week} тиждень"
    top_cell.font = title_font
    top_cell.fill = header_fill
    top_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Назви колонок
    headers = ["Пара", "Понеділок", "Вівторок", "Середа", "Четвер", "П'ятниця"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[2].height = 28

    items = ScheduleItem.objects.filter(week=current_week).select_related('teacher', 'subject', 'room')
    grid = {p: {} for p in range(1, 9)}
    for it in items:
        grid[it.period][it.day] = it

    days_keys = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    # Заповнення розкладу по рядках
    for p in range(1, 9):
        row_idx = p + 2
        # Збільшена висота рядка, щоб усе читалося комфортно
        ws.row_dimensions[row_idx].height = 65

        num_cell = ws.cell(row=row_idx, column=1, value=f"{p} пара")
        num_cell.font = bold_cell_font
        num_cell.fill = gray_fill
        num_cell.alignment = Alignment(horizontal="center", vertical="center")
        num_cell.border = thin_border

        for col_idx, d in enumerate(days_keys, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            if d == "Monday" and p == 8:
                cell.value = "—"
                cell.fill = gray_fill
                cell.font = cell_font
            else:
                item = grid[p].get(d)
                if item:
                    room_str = f" [ауд. {item.room.number}]" if item.room else ""
                    
                    # Чітке відображення групи або потоку
                    if item.is_stream:
                        group_line = f"★ ПОТІК: {item.groups}"
                        cell.fill = stream_fill
                    else:
                        group_line = f"Група: {item.groups}" if item.groups else ""
                        cell.fill = card_fill

                    # Формуємо текст комірки
                    parts = [item.subject.title, f"{item.teacher.name}{room_str}"]
                    if group_line:
                        parts.append(group_line)

                    cell.value = "\n".join(parts)
                    cell.font = cell_font
                else:
                    cell.value = ""

    # Ширина стовпців для охайного відображення
    ws.column_dimensions['A'].width = 14
    for c in ['B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[c].width = 34

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="schedule_week_{current_week}.xlsx"'
    wb.save(response)
    return response

def import_schedule_view(request):
    if request.method == "POST":
        file = request.FILES.get("schedule_file")
        if not file:
            messages.error(request, "Оберіть файл для завантаження!")
            return redirect("import_schedule")

        filename = file.name.lower()
        rows_to_process = []

        try:
            if filename.endswith(".xlsx"):
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if any(row):
                        rows_to_process.append([str(c).strip() if c is not None else "" for c in row[:7]])

            elif filename.endswith(".csv"):
                decoded_file = file.read().decode('utf-8-sig')
                io_string = io.StringIO(decoded_file)
                reader = csv.reader(io_string, delimiter=';' if ';' in decoded_file[:200] else ',')
                next(reader, None)
                for row in reader:
                    if any(row):
                        rows_to_process.append([c.strip() for c in row[:7]])
            else:
                messages.error(request, "Формат не підтримується! Завантажте .xlsx або .csv")
                return redirect("import_schedule")

            imported_count = 0
            for r in rows_to_process:
                if len(r) < 5:
                    continue
                week_str, day_raw, period_str, subj_title, teacher_name = r[0], r[1], r[2], r[3], r[4]
                groups = r[5] if len(r) > 5 and r[5] else ""
                room_num = r[6] if len(r) > 6 and r[6] else None

                day_map = {'monday':'Monday','понеділок':'Monday','tuesday':'Tuesday','вівторок':'Tuesday','wednesday':'Wednesday','середа':'Wednesday','thursday':'Thursday','четвер':'Thursday','friday':'Friday',"п'ятниця":'Friday'}
                day = day_map.get(day_raw.lower())
                if not day: continue
                period = int(period_str)
                if day == "Monday" and period > 7: continue

                teacher, _ = Teacher.objects.get_or_create(name=teacher_name)
                subject, _ = Subject.objects.get_or_create(title=subj_title)
                room = Room.objects.get_or_create(number=room_num)[0] if room_num else None

                ScheduleItem.objects.update_or_create(
                    week=int(week_str),
                    day=day,
                    period=period,
                    defaults={'teacher': teacher, 'subject': subject, 'room': room, 'groups': groups}
                )
                imported_count += 1

            if imported_count > 0:
                messages.success(request, f"Успішно імпортовано {imported_count} пар!")
        except Exception as e:
            messages.error(request, f"Помилка: {str(e)}")

        return redirect("import_schedule")

    return render(request, "import.html")
